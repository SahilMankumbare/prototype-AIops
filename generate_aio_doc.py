from docx import Document

doc = Document()
doc.add_heading('AIOps System Technical Overview', 0)
doc.add_paragraph('This document explains each script in the AIOps repository, outlines the role of the Ollama AI API, describes where logs are sourced from, and justifies why this solution is more than simple automation.')

doc.add_heading('1. config.py', level=1)
doc.add_paragraph('Purpose: Holds device configuration and system settings for the AIOps environment.')
for item in [
    'Defines the list of managed devices with SSH/login credentials and device types.',
    'Specifies polling interval, dashboard port, and AI provider URL/model.',
    'Acts as a single source of truth for all runtime connectivity settings.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('2. collector.py', level=1)
doc.add_paragraph('Purpose: Polls Cisco devices over SSH and builds structured operational data for monitoring and AI analysis.')
for item in [
    'Uses Netmiko to connect to each configured device.',
    'Runs `show interfaces` and parses the output into per-interface objects with status, link state, and protocol.',
    'Runs `show access-lists` and parses ACL entries into structured rule dictionaries.',
    'Detects unreachable devices, timeouts, and authentication failures.',
    'Provides helper functions such as `get_down_interfaces()` to expose interface failures to the AI agent.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('3. router_api.py', level=1)
doc.add_paragraph('Purpose: Hosts the Flask dashboard API, serves the UI, accepts manual commands, and handles log retrieval and PDF export.')
for item in [
    'Runs a Flask web server with endpoints for device status and logs.',
    'Starts a background poller to continuously refresh current device state by calling `collector.poll_all_devices()`.',
    'Exposes `/execute-command` for natural-language manual commands and `/download-logs-pdf` for exporting router logs.',
    'Maintains in-memory state for device status, runtime logs, and fetched router logs.',
    'Uses AI-assisted parsing to convert user commands into structured actions before executing them on the router.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph('Special log-related behavior:')
for item in [
    'The `show_logs` action connects to the router and executes actual IOS commands such as `show logging`, `show history`, `show ip int brief`, `show interfaces status`, `show access-lists`, `show ip ssh`, and `show ip traffic`.',
    'Command output is stored in `state["router_logs"]` and exported as a sectioned PDF.',
    'The PDF generator formats each command output block as a separate section heading to make the document readable and organized.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('4. ai_agent.py', level=1)
doc.add_paragraph('Purpose: Implements the AI observation/response loop and decides whether to auto-fix or escalate network issues.')
for item in [
    'Polls all devices using `collector.poll_all_devices()`.',
    'Builds a natural-language prompt containing live device status and down-interface details.',
    'Sends that prompt to the Ollama API and expects structured JSON in return.',
    'Parses the AI response to decide whether each interface problem should be AUTO-FIXed or ESCALATEd.',
    'Calls remediation code when the AI decides to auto-fix a failed interface.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('5. remediation.py', level=1)
doc.add_paragraph('Purpose: Applies corrective configuration changes to devices when the AI decides remediation is needed.')
for item in [
    'Looks up device credentials from config.py.',
    'Connects over SSH to the target device.',
    'Sends configuration commands such as `interface <X>` and `no shutdown` to restore down interfaces.',
    'Saves the device configuration and reports success or failure.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('6. drift_detector.py', level=1)
doc.add_paragraph('Purpose: Uses AI to analyze configuration diffs and classify configuration drift risks.')
for item in [
    'Accepts a configuration diff string and device name.',
    'Builds a prompt asking Ollama to classify the diff as BENIGN, RISKY, or CRITICAL.',
    'Requests AI-generated rollback commands and a risk summary.',
    'Returns structured metadata to help decide whether a config change needs remediation or rollback.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('7. dashboard.html', level=1)
doc.add_paragraph('Purpose: Provides the front-end UI for live device status, manual commands, and log access.')
for item in [
    'Fetches device status and logs from Flask endpoints at regular intervals.',
    'Allows the user to enter freeform commands and send them to `/execute-command`.',
    'Displays command results, current topology status, and a logs panel.',
    'Includes a manual PDF download button for the latest router logs.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('AI / Ollama Role in the System', level=1)
doc.add_paragraph('Ollama is the local AI model server that provides multiple intelligence layers in this system.')
for item in [
    'Natural-language command parsing in `router_api.py`.',
    'Network issue analysis and remediation decisions in `ai_agent.py`.',
    'Configuration drift classification in `drift_detector.py`.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph('How Ollama is used:')
doc.add_paragraph('The repository sends prompts to `http://localhost:11434/api/generate` using `AI_MODEL`. Ollama returns generated text, which the code parses into JSON for structured workflows.')

doc.add_heading('Where logs come from', level=1)
doc.add_paragraph('The system retrieves logs directly from the router devices themselves, not from local application logs.')
for item in [
    'Manual log requests execute actual IOS commands on the router via SSH in `router_api.py`.',
    'Periodic polling in `collector.py` reads interface and ACL state from devices.',
    'The PDF report is built from the router command outputs and organized into command sections.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Why this is more than automation', level=1)
doc.add_paragraph('This system is not only scripting; it blends monitoring, AI analysis, and conditional remediation:')
for item in [
    'AI-powered reasoning to decide whether a problem should be auto-fixed or escalated.',
    'Natural-language command understanding and structured action parsing.',
    'Live network polling plus manual command control.',
    'Sectioned PDF reports from actual router outputs.',
    'Configuration drift analysis and rollback guidance through AI models.'
]:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Conclusion', level=1)
doc.add_paragraph('The repository is designed as a practical AIOps toolkit that combines real network device access, AI reasoning, and structured operational output. It fetches logs from live routers, uses Ollama for intelligence, and produces actionable results beyond simple task automation.')

doc.save('AIOps_System_Overview.docx')
print('AIOps_System_Overview.docx created')
